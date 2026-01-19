import streamlit as st
import streamlit.components.v1 as components

import google.generativeai as genai
from PIL import Image, ImageEnhance, ImageOps, ImageFilter
import pypdfium2 as pdfium

import io, gc, base64, time, random, html, re
from datetime import datetime
from collections import Counter

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from supabase import create_client


# ==========================================
# 1) CONFIG
# ==========================================
st.set_page_config(
    page_title="Manuscript AI - Open Academic Portal",
    page_icon="📜",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ==========================================
# 2) APP CONSTANTS
# ==========================================
THEME = "DARK_GOLD"
DEMO_LIMIT_PAGES = 3
STARTER_CREDITS = 10
HISTORY_LIMIT = 20
BATCH_DELAY_RANGE = (0.8, 1.6)
MAX_OUT_TOKENS = 4096

# ==========================================
# 3) THEMES
# ==========================================
THEMES = {
    "DARK_GOLD": {
        "app_bg": "#0b1220",
        "surface": "#10182b",
        "sidebar_bg": "#0c1421",
        "text": "#eaf0ff",
        "muted": "#c7d0e6",
        "gold": "#c5a059",
        "gold2": "#d4af37",
    }
}
C = THEMES["DARK_GOLD"]

# ==========================================
# 4) CSS
# ==========================================
st.markdown(f"""
<style>
:root {{
  --app-bg: {C["app_bg"]};
  --surface: {C["surface"]};
  --sidebar-bg: {C["sidebar_bg"]};
  --text: {C["text"]};
  --muted: {C["muted"]};
  --gold: {C["gold"]};
  --gold2: {C["gold2"]};
}}

html, body {{
  background: var(--app-bg) !important;
  margin: 0 !important;
  padding: 0 !important;
}}

.stApp, div[data-testid="stAppViewContainer"] {{
  background: var(--app-bg) !important;
  min-height: 100vh !important;
}}

div[data-testid="stAppViewContainer"] .main .block-container {{
  padding-top: 3.25rem !important;
  padding-bottom: 1.25rem !important;
}}

footer {{visibility: hidden !important;}}
.stAppDeployButton {{display:none !important;}}
#stDecoration {{display:none !important;}}
header[data-testid="stHeader"] {{ background: rgba(0,0,0,0) !important; }}

section[data-testid="stSidebar"] {{
  background: var(--sidebar-bg) !important;
  border-right: 2px solid var(--gold) !important;
}}
section[data-testid="stSidebar"] * {{
  color: var(--text) !important;
}}
section[data-testid="stSidebar"] .stCaption {{
  color: var(--muted) !important;
}}

h1, h2, h3, h4 {{
  color: var(--gold) !important;
  font-family: 'Georgia', serif;
  border-bottom: 2px solid var(--gold) !important;
  padding-bottom: 8px !important;
  text-align: center !important;
  text-shadow: 0 1px 1px rgba(0,0,0,0.35);
}}

.stMarkdown p {{
  color: var(--muted) !important;
}}

.stButton>button {{
  background: linear-gradient(135deg, var(--sidebar-bg) 0%, #1e3a8a 100%) !important;
  color: var(--gold) !important;
  font-weight: 900 !important;
  width: 100% !important;
  padding: 11px 12px !important;
  border: 1px solid var(--gold) !important;
  border-radius: 12px !important;
  box-shadow: 0 10px 22px rgba(0,0,0,0.25) !important;
  transition: transform .15s ease, filter .2s ease, box-shadow .2s ease !important;
}}
.stButton>button:hover {{
  transform: translateY(-1px);
  filter: brightness(1.08);
  box-shadow: 0 14px 28px rgba(0,0,0,0.32) !important;
}}

.stTextInput input, .stSelectbox select {{
  background-color: rgba(255,255,255,0.06) !important;
  color: var(--text) !important;
  border: 1px solid rgba(197,160,89,0.55) !important;
  border-radius: 10px !important;
}}

.stTextArea textarea {{
  background-color: #fdfaf1 !important;
  color: #000000 !important;
  border: 1px solid rgba(197,160,89,0.55) !important;
  border-radius: 10px !important;
}}

.chat-user {{
  background-color: #e2e8f0;
  color: #000;
  padding: 10px;
  border-radius: 10px;
  border-left: 5px solid #1e3a8a;
  margin-bottom: 6px;
}}
.chat-ai {{
  background-color: #ffffff;
  color: #1a1a1a;
  padding: 10px;
  border-radius: 10px;
  border: 1px solid #d4af37;
  margin-bottom: 14px;
}}

.sticky-preview {{
  position: sticky;
  top: 4.6rem;
  border-radius: 14px;
  border: 2px solid var(--gold);
  overflow: hidden;
  box-shadow: 0 14px 35px rgba(0,0,0,0.22);
  background: rgba(0,0,0,0.15);
  max-height: 540px;
  animation: fadeUp .35s ease both;
}}
.sticky-preview img {{
  width: 100%;
  height: 540px;
  object-fit: contain;
  display: block;
}}

@keyframes fadeUp {{
  from {{ opacity: 0; transform: translateY(10px); }}
  to   {{ opacity: 1; transform: translateY(0); }}
}}

@media (max-width: 768px) {{
  div[data-testid="stAppViewContainer"] .main .block-container {{
    padding-top: 3.5rem !important;
    padding-left: 1rem !important;
    padding-right: 1rem !important;
  }}
  .sticky-preview {{ position: relative; top: 0; max-height: 48vh; }}
  .sticky-preview img {{ height: 48vh; }}
}}
</style>
""", unsafe_allow_html=True)

# ==========================================
# 5) SERVICES
# ==========================================
@st.cache_resource
def get_db():
    return create_client(st.secrets["SUPABASE_URL"], st.secrets["SUPABASE_KEY"])

db = get_db()

genai.configure(api_key=st.secrets["GEMINI_API_KEY"])
model = genai.GenerativeModel(model_name="gemini-flash-latest")  # o'zgarmaydi

# ==========================================
# 6) STATE
# ==========================================
if "auth" not in st.session_state: st.session_state.auth = False
if "u_email" not in st.session_state: st.session_state.u_email = ""
if "last_fn" not in st.session_state: st.session_state.last_fn = None

if "page_bytes" not in st.session_state: st.session_state.page_bytes = []
if "results" not in st.session_state: st.session_state.results = {}
if "chats" not in st.session_state: st.session_state.chats = {}
if "warn_db" not in st.session_state: st.session_state.warn_db = False

# ==========================================
# 7) HELPERS
# ==========================================
def pil_to_jpeg_bytes(img: Image.Image, quality: int = 90, max_side: int = 2800) -> bytes:
    img = img.convert("RGB")
    w, h = img.size
    long_side = max(w, h)
    if long_side > max_side:
        ratio = max_side / float(long_side)
        img = img.resize((int(w * ratio), int(h * ratio)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=quality, optimize=True)
    return buf.getvalue()

@st.cache_data(show_spinner=False, max_entries=12)
def render_pdf_pages_to_bytes(file_bytes: bytes, max_pages: int, scale: float) -> list[bytes]:
    pdf = pdfium.PdfDocument(file_bytes)
    out: list[bytes] = []
    try:
        n = min(len(pdf), max_pages)
        for i in range(n):
            pil_img = pdf[i].render(scale=scale).to_pil()
            out.append(pil_to_jpeg_bytes(pil_img, quality=88, max_side=3000))
    finally:
        try: pdf.close()
        except Exception: pass
    return out

@st.cache_data(show_spinner=False, max_entries=256)
def preprocess_bytes(img_bytes: bytes, brightness: float, contrast: float, rotate: int, sharpen: float) -> bytes:
    img = Image.open(io.BytesIO(img_bytes))
    img = ImageOps.exif_transpose(img)
    if rotate:
        img = img.rotate(rotate, expand=True)
    img = ImageEnhance.Brightness(img).enhance(brightness)
    img = ImageEnhance.Contrast(img).enhance(contrast)
    if sharpen > 0:
        img = img.filter(ImageFilter.UnsharpMask(radius=2, percent=int(120 * sharpen), threshold=2))
    return pil_to_jpeg_bytes(img, quality=90, max_side=2800)

def parse_pages(spec: str, max_n: int) -> list[int]:
    spec = (spec or "").strip()
    if not spec:
        return [0] if max_n > 0 else []
    out = set()
    parts = [p.strip() for p in spec.split(",") if p.strip()]
    for part in parts:
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                a = int(a.strip()); b = int(b.strip())
                if a > b: a, b = b, a
                for p in range(a, b + 1):
                    if 1 <= p <= max_n:
                        out.add(p - 1)
            else:
                p = int(part)
                if 1 <= p <= max_n:
                    out.add(p - 1)
        except Exception:
            continue
    return sorted(out) if out else ([0] if max_n > 0 else [])

def call_gemini_with_retry(prompt: str, payloads: list[dict], tries: int = 3) -> str:
    last_err = None
    for i in range(tries):
        try:
            resp = model.generate_content(
                [prompt, *payloads],
                generation_config={
                    "max_output_tokens": MAX_OUT_TOKENS,
                    "temperature": 0.2,
                }
            )
            return getattr(resp, "text", "") or ""
        except Exception as e:
            last_err = e
            msg = str(e).lower()
            if ("429" in msg) or ("rate" in msg) or ("quota" in msg) or ("resource" in msg):
                time.sleep((2 ** i) + random.random())
                continue
            raise
    raise RuntimeError("So'rovlar ko'p (429). Birozdan keyin qayta urinib ko'ring.") from last_err

def ensure_profile(email: str) -> None:
    try:
        existing = db.table("profiles").select("email,credits").eq("email", email).limit(1).execute()
        if existing.data:
            return
        db.table("profiles").insert({"email": email, "credits": STARTER_CREDITS}).execute()
    except Exception:
        st.session_state.warn_db = True

def get_credits(email: str) -> int:
    try:
        r = db.table("profiles").select("credits").eq("email", email).single().execute()
        return int(r.data["credits"]) if r.data and "credits" in r.data else 0
    except Exception:
        st.session_state.warn_db = True
        return 0

def consume_credit_safe(email: str, n: int = 1) -> bool:
    try:
        r = db.rpc("consume_credits", {"p_email": email, "p_n": n}).execute()
        return bool(r.data)
    except Exception:
        pass
    for _ in range(2):
        try:
            cur = get_credits(email)
            if cur < n:
                return False
            newv = cur - n
            upd = db.table("profiles").update({"credits": newv}).eq("email", email).eq("credits", cur).execute()
            if upd.data:
                return True
        except Exception:
            st.session_state.warn_db = True
            return False
    return False

def refund_credit_safe(email: str, n: int = 1) -> None:
    try:
        db.rpc("refund_credits", {"p_email": email, "p_n": n}).execute()
        return
    except Exception:
        pass
    for _ in range(2):
        try:
            cur = get_credits(email)
            upd = db.table("profiles").update({"credits": cur + n}).eq("email", email).eq("credits", cur).execute()
            if upd.data:
                return
        except Exception:
            st.session_state.warn_db = True
            return

def log_usage(email: str, doc_name: str, page_index: int, status: str, note: str = "") -> None:
    try:
        db.table("usage_logs").insert({
            "email": email,
            "doc_name": doc_name,
            "page_index": page_index,
            "status": status,
            "note": note[:240],
            "created_at": datetime.utcnow().isoformat()
        }).execute()
    except Exception:
        st.session_state.warn_db = True

def save_report(email: str, doc_name: str, page_index: int, result_text: str) -> None:
    try:
        db.table("reports").upsert(
            {
                "email": email,
                "doc_name": doc_name,
                "page_index": page_index,
                "result_text": result_text,
                "updated_at": datetime.utcnow().isoformat()
            },
            on_conflict="email,doc_name,page_index"
        ).execute()
    except Exception:
        try:
            db.table("reports").insert({
                "email": email,
                "doc_name": doc_name,
                "page_index": page_index,
                "result_text": result_text,
                "created_at": datetime.utcnow().isoformat()
            }).execute()
        except Exception:
            st.session_state.warn_db = True

def load_reports(email: str, doc_name: str) -> dict:
    try:
        r = db.table("reports").select("page_index,result_text").eq("email", email).eq("doc_name", doc_name).limit(200).execute()
        out = {}
        for row in (r.data or []):
            out[int(row["page_index"])] = row.get("result_text") or ""
        return out
    except Exception:
        st.session_state.warn_db = True
        return {}

def extract_diagnosis(text: str) -> dict:
    t = text or ""
    def pick(rx):
        m = re.search(rx, t, flags=re.IGNORECASE)
        return (m.group(1).strip() if m else "").strip()
    til = pick(r"Til\s*:\s*(.+)")
    xat = pick(r"Xat\s*uslubi\s*:\s*(.+)")
    conf = pick(r"Ishonchlilik\s*:\s*(.+)")
    return {"til": til.replace("|", "").strip(),
            "xat": xat.replace("|", "").strip(),
            "conf": conf.replace("|", "").strip()}

def aggregate_detected_meta(results: dict[int, str]) -> dict:
    til_list, xat_list, conf_list = [], [], []
    for _, txt in results.items():
        d = extract_diagnosis(txt)
        if d["til"] and "noma" not in d["til"].lower(): til_list.append(d["til"])
        if d["xat"] and "noma" not in d["xat"].lower(): xat_list.append(d["xat"])
        if d["conf"]: conf_list.append(d["conf"])
    til = Counter(til_list).most_common(1)[0][0] if til_list else ""
    xat = Counter(xat_list).most_common(1)[0][0] if xat_list else ""
    conf = Counter(conf_list).most_common(1)[0][0] if conf_list else ""
    return {"til": til, "xat": xat, "conf": conf}


# ============================
# ANALYSIS QUALITY PATCH (Tez / To‘liq) - FAQAT TAHLIL LOGIKASI
# ============================
OCR_OVERLAP = 0.12

def get_analysis_cfg(mode: str) -> dict:
    mode = (mode or "Tez").strip()
    if mode == "To‘liq":
        return {
            "min_chars": 650,
            "strips_default": 10,
            "strips_retry": 13,
            "strip_sleep": (0.18, 0.38),
            "max_calls": 30,
            "chunk_chars": 2200,
            "max_chunks": 8,
        }
    return {
        "min_chars": 420,
        "strips_default": 6,
        "strips_retry": 9,
        "strip_sleep": (0.12, 0.28),
        "max_calls": 18,
        "chunk_chars": 3200,
        "max_chunks": 5,
    }

def _safe_model_call_budget_guard(counter: dict, add: int = 1, max_calls: int = 18):
    counter["n"] = int(counter.get("n", 0)) + int(add)
    if counter["n"] > int(max_calls):
        raise RuntimeError("Sahifa bo‘yicha AI chaqiruv limiti oshdi (xavfsizlik).")

def _call_gemini_guarded(prompt: str, payloads: list[dict], call_budget: dict, cfg: dict, tries: int = 3) -> str:
    _safe_model_call_budget_guard(call_budget, 1, max_calls=int(cfg.get("max_calls", 18)))
    return call_gemini_with_retry(prompt, payloads, tries=tries)

def _split_two_pages_if_wide(img: Image.Image) -> list[Image.Image]:
    w, h = img.size
    if w >= int(h * 1.15):
        mid = w // 2
        return [img.crop((0, 0, mid, h)), img.crop((mid, 0, w, h))]
    return [img]

def _split_vertical_strips(img: Image.Image, n: int, overlap: float) -> list[Image.Image]:
    w, h = img.size
    step = h / float(max(n, 1))
    ov = int(step * overlap)
    strips: list[Image.Image] = []
    for i in range(max(n, 1)):
        y0 = max(0, int(i * step) - ov)
        y1 = min(h, int((i + 1) * step) + ov)
        strips.append(img.crop((0, y0, w, y1)))
    return strips

def _jpeg_payload_from_pil(img: Image.Image) -> dict:
    jb = pil_to_jpeg_bytes(img, quality=90, max_side=3000)
    return {"mime_type": "image/jpeg", "data": base64.b64encode(jb).decode("utf-8")}

def has_read_sections(text: str) -> bool:
    if not text:
        return False
    t = text.lower()
    return ("0) tashxis" in t) and ("1) transliteratsiya" in t)

def has_final_sections(text: str) -> bool:
    if not text:
        return False
    must = [
        "0) Tashxis",
        "1) Transliteratsiya",
        "2) To‘g‘ridan-to‘g‘ri tarjima",
        "3) Akademik tarjima",
        "4) Paleografiya",
        "5) Arxaik",
        "6) Izoh",
    ]
    t = text.lower()
    return all(m.lower() in t for m in must)

def extract_diag_block(read_text: str) -> str:
    if not read_text:
        return ""
    m = re.search(r"(0\)\s*Tashxis[\s\S]+?)(?=\n\s*1\)\s*Transliteratsiya)", read_text, flags=re.IGNORECASE)
    return (m.group(1).strip() if m else "").strip()

def extract_translit_lines(read_text: str) -> str:
    if not read_text:
        return ""
    m = re.search(r"1\)\s*Transliteratsiya[^\n]*\n([\s\S]+)", read_text, flags=re.IGNORECASE)
    if not m:
        return ""
    return m.group(1).strip()

def build_read_text(diag_block: str, translit_lines: str) -> str:
    diag = (diag_block or "").strip()
    tl = (translit_lines or "").strip()
    if not diag:
        diag = (
            "0) Tashxis:\n"
            "Til: Noma'lum\n"
            "Xat uslubi: Noma'lum\n"
            "Hint mosligi: hint yo‘q\n"
            "Ishonchlilik: o‘rtacha\n"
            "Sabab: OCR/rasm sifati va model taxmini.\n"
        )
    return f"{diag}\n\n1) Transliteratsiya (satrma-satr, to‘liq):\n{tl}".strip()

def translit_min_len_ok(translit_lines: str, min_chars: int = 600) -> bool:
    return bool(translit_lines and len(translit_lines.strip()) >= int(min_chars))

def read_transliteration_with_strips(p_read: str, img_bytes: bytes, strips_n: int, min_chars: int, call_budget: dict, cfg: dict) -> tuple[str, str]:
    try:
        payload_full = {"mime_type": "image/jpeg", "data": base64.b64encode(img_bytes).decode("utf-8")}
        read0 = _call_gemini_guarded(p_read, [payload_full], call_budget, cfg, tries=2).strip()
        diag_block = extract_diag_block(read0)

        pil = Image.open(io.BytesIO(img_bytes)).convert("RGB")
        pages = _split_two_pages_if_wide(pil)

        all_lines: list[str] = []
        for p_i, p in enumerate(pages, start=1):
            strips = _split_vertical_strips(p, n=int(strips_n), overlap=OCR_OVERLAP)
            for s_i, strip in enumerate(strips, start=1):
                payload = _jpeg_payload_from_pil(strip)
                strip_prompt = (
                    p_read
                    + "\n\nMUHIM: Bu rasm BO‘LAK. Natijada faqat shu bo‘lakdagi matnning TRANSLITERATSIYA SATRLARINI yozing.\n"
                      "- Hech qanday sarlavha (0)/1) yozmang.\n"
                      "- Hech qanday izoh yozmang.\n"
                      "- Har satr alohida qatorda.\n"
                      "- O‘qilmasa: [o‘qilmadi] yoki [?].\n"
                    + f"\n(Bo‘lak: bet {p_i}, {s_i}/{len(strips)})"
                )
                txt = _call_gemini_guarded(strip_prompt, [payload], call_budget, cfg, tries=2).strip()
                if txt:
                    all_lines.append(txt)
                time.sleep(random.uniform(*cfg.get("strip_sleep", (0.12, 0.28))))

        translit_lines = "\n".join([x for x in all_lines if x.strip()]).strip()
        if not translit_min_len_ok(translit_lines, min_chars=min_chars):
            return diag_block, translit_lines
        return diag_block, translit_lines
    except Exception:
        return "", ""

def chunk_text_safe(text: str, chunk_chars: int = 2200, overlap: int = 120, max_chunks: int | None = None) -> list[str]:
    t = (text or "").strip()
    if len(t) <= chunk_chars:
        return [t]
    out: list[str] = []
    i = 0
    while i < len(t):
        j = min(len(t), i + chunk_chars)
        out.append(t[i:j])
        if max_chunks is not None and len(out) >= int(max_chunks):
            break
        if j >= len(t):
            break
        i = max(0, j - overlap)
    return out

def _extract_tag(resp: str, tag: str) -> str:
    if not resp:
        return ""
    m = re.search(rf"<{tag}>([\s\S]+?)</{tag}>", resp, flags=re.IGNORECASE)
    return (m.group(1).strip() if m else "").strip()

def translate_2_3_chunked(translit_lines: str, call_budget: dict, cfg: dict) -> tuple[str, str]:
    chunks = chunk_text_safe(
        translit_lines,
        chunk_chars=int(cfg.get("chunk_chars", 3200)),
        overlap=160,
        max_chunks=int(cfg.get("max_chunks", 5))
    )
    direct_parts: list[str] = []
    acad_parts: list[str] = []

    for k, ch in enumerate(chunks, start=1):
        prompt = (
            "Siz qadimiy qo‘lyozmalar tarjimonisiz.\n"
            "Vazifa: berilgan TRANSLITERATSIYA BO‘LAGI asosida 2 xil tarjima qiling.\n"
            "QOIDALAR:\n"
            "- Qisqartirmang.\n"
            "- Har satrni imkon qadar satrma-satr tarjima qiling.\n"
            "- O‘qilmagan joy: [o‘qilmadi] yoki [?] ni saqlang.\n"
            "- Natija faqat o‘zbekcha.\n\n"
            "FAqat quyidagi formatda chiqaring (hech narsa qo‘shmang):\n"
            "<D>\n(2) To‘g‘ridan-to‘g‘ri tarjima satrlari\n</D>\n"
            "<A>\n(3) Akademik tarjima satrlari\n</A>\n\n"
            f"(Bo‘lak {k}/{len(chunks)})\n"
            "TRANSLITERATSIYA:\n"
            f"{ch}\n"
        )
        resp = _call_gemini_guarded(prompt, [], call_budget, cfg, tries=3).strip()
        d = _extract_tag(resp, "D") or resp
        a = _extract_tag(resp, "A") or ""
        direct_parts.append(d.strip())
        if a.strip():
            acad_parts.append(a.strip())
        time.sleep(random.uniform(0.18, 0.45))

    direct = "\n".join([x for x in direct_parts if x.strip()]).strip()
    acad = "\n".join([x for x in acad_parts if x.strip()]).strip()
    return direct, acad

def analyze_4_5_6_once(p_an: str, translit_lines: str, call_budget: dict, cfg: dict) -> tuple[str, str, str]:
    excerpt = translit_lines.strip()
    if len(excerpt) > 9000:
        excerpt = excerpt[:9000] + "\n... [davomi qisqartirildi: tahlil uchun excerpt] ..."

    prompt = (
        p_an
        + "\n\nMUHIM: Faqat 4), 5), 6) bo‘limlar mazmunini chiqaring. 2) va 3) ni yozmang.\n"
          "Natija faqat quyidagi TAG formatda bo‘lsin (hech narsa qo‘shmang):\n"
          "<P>\n4) Paleografiya matni\n</P>\n"
          "<L>\n5) Arxaik lug'at (5–10 so‘z)\n</L>\n"
          "<I>\n6) Izoh\n</I>\n\n"
          "TRANSLITERATSIYA (asos):\n"
        + excerpt
    )

    resp = _call_gemini_guarded(prompt, [], call_budget, cfg, tries=3).strip()
    p = _extract_tag(resp, "P") or ""
    l = _extract_tag(resp, "L") or ""
    i = _extract_tag(resp, "I") or ""
    return p.strip(), l.strip(), i.strip()

def build_analyze_text(p_an: str, translit_lines: str, call_budget: dict, cfg: dict) -> str:
    direct, acad = translate_2_3_chunked(translit_lines, call_budget, cfg)

    if not acad.strip():
        prompt_retry = (
            "Siz Manuscript AI tarjimonisiz.\n"
            "Vazifa: quyidagi transliteratsiya asosida AKADEMIK tarjima yozing.\n"
            "Qoidalar: qisqartirmang, o‘zbekcha, izchil.\n\n"
            "FAqat shu formatda:\n"
            "3) Akademik tarjima (mazmuniy, izchil):\n"
            "<matn>\n\n"
            "TRANSLITERATSIYA:\n"
            + (translit_lines[:4500] if len(translit_lines) > 4500 else translit_lines)
        )
        acad = _call_gemini_guarded(prompt_retry, [], call_budget, cfg, tries=2).strip()

    paleo, lex, note = analyze_4_5_6_once(p_an, translit_lines, call_budget, cfg)

    if not direct.strip():
        direct = "[Tarjima chiqmadi]"
    if not paleo.strip():
        paleo = "[Paleografiya bo‘limi chiqmadi]"
    if not lex.strip():
        lex = "[Arxaik lug‘at chiqmadi]"
    if not note.strip():
        note = "[Izoh chiqmadi]"

    out = (
        "2) To‘g‘ridan-to‘g‘ri tarjima (oddiy o‘zbekcha):\n"
        f"{direct.strip()}\n\n"
        "3) Akademik tarjima (mazmuniy, izchil):\n"
        f"{acad.strip()}\n\n"
        "4) Paleografiya:\n"
        f"{paleo.strip()}\n\n"
        "5) Arxaik lug'at (5–10 so‘z):\n"
        f"{lex.strip()}\n\n"
        "6) Izoh (kontekst; aniq bo‘lmasa ehtiyot bo‘l):\n"
        f"{note.strip()}\n"
    )
    return out.strip()


# ==========================================
# PROMPTS
# ==========================================
def build_prompts(hint_lang: str, hint_era: str) -> tuple[str, str]:
    p_read = (
        "Siz qo‘lyozma o‘qish bo‘yicha mutaxassissiz.\n"
        "Vazifa: rasm ichidagi yozuvni maksimal to‘liq o‘qing.\n"
        "QOIDALAR:\n"
        "- Hech narsani qisqartirmang.\n"
        "- Taxmin qilmang: o‘qilmasa [o‘qilmadi] yoki [?] yozing.\n"
        "- Ism/son/sana/joylarda faqat ko‘ringanini yozing.\n"
        "- Natija faqat o‘zbekcha bo‘lsin.\n\n"
        f"Foydalanuvchi hintlari: til='{hint_lang or 'yo‘q'}', xat='{hint_era or 'yo‘q'}' (faqat hint).\n\n"
        "FORMAT (aniq shunday):\n"
        "0) Tashxis:\n"
        "Til: <aniqlangan til>\n"
        "Xat uslubi: <aniqlangan xat uslubi>\n"
        "Hint mosligi: <mos | mos emas | hint yo‘q>\n"
        "Ishonchlilik: past | o‘rtacha | yuqori\n"
        "Sabab: <1-2 jumla>\n"
        "1) Transliteratsiya (satrma-satr, to‘liq):\n"
        "- Har satrni alohida qatorga yozing.\n"
        "- Matnda bo‘lim/ustun bo‘lsa, '--- USTUN 1 ---' kabi ajrating.\n"
    )

    p_an = (
        "Siz Manuscript AI mutaxassisiz.\n"
        "Vazifa: berilgan transliteratsiya asosida tarjima va akademik tahlil qiling.\n"
        "QOIDALAR:\n"
        "- Hech narsani uydirmang.\n"
        "- Ism/son/sana/joylarni transliteratsiyadagi ko‘rinishiga qat’iy mos yozing.\n"
        "- Natija faqat o‘zbekcha.\n\n"
        "FORMAT (aniq shunday):\n"
        "2) To‘g‘ridan-to‘g‘ri tarjima (oddiy o‘zbekcha):\n"
        "3) Akademik tarjima (mazmuniy, izchil):\n"
        "4) Paleografiya:\n"
        "5) Arxaik lug'at (5–10 so‘z):\n"
        "6) Izoh (kontekst; aniq bo‘lmasa ehtiyot bo‘l):\n"
    )
    return p_read, p_an


# ==========================================
# 8) RESULT CARD HTML
# ==========================================
def _badge(conf: str) -> str:
    conf_l = (conf or "").lower()
    if "yuqori" in conf_l:
        cls = "b-high"; label = "Yuqori ishonch"
    elif "o‘rtacha" in conf_l or "ortacha" in conf_l:
        cls = "b-med"; label = "O‘rtacha ishonch"
    else:
        cls = "b-low"; label = "Past ishonch"
    return f'<span class="badge {cls}">{html.escape(label)}</span>'

def md_to_html(md: str) -> str:
    raw = md or ""
    diag = extract_diagnosis(raw)
    safe = html.escape(raw)
    lines = safe.splitlines()
    out = []
    i = 0

    def is_h(nline: str) -> bool:
        return bool(re.match(r"^\d+\)\s+", nline.strip()))

    while i < len(lines):
        line = lines[i].rstrip()

        if line.strip().lower().startswith("0)") and "tashxis" in line.strip().lower():
            block = [line]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("1)"):
                block.append(lines[i].rstrip())
                i += 1
            badge = _badge(diag.get("conf", ""))
            block_html = "<br/>".join(block)
            out.append(f"""
              <div class="diag-box">
                <div class="diag-head">
                  <span class="diag-title">0) Tashxis</span>
                  {badge}
                </div>
                <div class="diag-body">{block_html}</div>
              </div>
            """)
            continue

        if line.startswith("## "):
            out.append(f"<h2>{line[3:]}</h2>"); i += 1; continue
        if line.startswith("### "):
            out.append(f"<h3>{line[4:]}</h3>"); i += 1; continue
        if is_h(line):
            out.append(f"<h3>{line.strip()}</h3>"); i += 1; continue
        if line.strip() == "---":
            out.append("<hr/>"); i += 1; continue
        if line.strip() == "":
            out.append("<br/>"); i += 1; continue

        out.append(f"<p style='white-space:pre-wrap; margin:10px 0;'>{line}</p>")
        i += 1

    return "\n".join(out)

def render_result_card(md: str, gold: str) -> str:
    body = md_to_html(md)
    return f"""
    <style>
      :root {{ --gold: {gold}; }}
      body {{ margin:0; font-family: Georgia, 'Times New Roman', serif; color:#111827; background:#ffffff; }}
      .card {{
        padding: 18px;
        border-left: 10px solid var(--gold);
        border-radius: 16px;
        box-shadow: 0 10px 30px rgba(0,0,0,0.18);
        line-height: 1.75;
      }}
      h2, h3 {{
        margin: 0 0 10px 0;
        color: var(--gold);
        border-bottom: 2px solid var(--gold);
        padding-bottom: 8px;
      }}
      .diag-box {{
        border: 1px solid rgba(17, 24, 39, 0.10);
        background: linear-gradient(180deg, rgba(197,160,89,0.14), rgba(197,160,89,0.06));
        border-radius: 14px;
        padding: 12px 12px;
        margin: 6px 0 14px 0;
      }}
      .diag-head {{ display:flex; align-items:center; justify-content:space-between; gap:10px; margin-bottom:8px; }}
      .diag-title {{ font-weight: 900; color: #1f2937; }}
      .badge {{
        font-size: 12px; font-weight: 900;
        padding: 6px 10px; border-radius: 999px;
        border: 1px solid rgba(0,0,0,0.08);
        box-shadow: 0 6px 14px rgba(0,0,0,0.08);
        white-space: nowrap;
      }}
      .b-high {{ background: #dcfce7; color: #14532d; }}
      .b-med  {{ background: #fef9c3; color: #713f12; }}
      .b-low  {{ background: #fee2e2; color: #7f1d1d; }}
      hr {{ border: none; border-top: 1px solid #e5e7eb; margin: 14px 0; }}
    </style>
    <div class="card">{body}</div>
    """

# ==========================================
# 9) WORD EXPORT
# ==========================================
def _doc_set_normal_style(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

def _add_cover(doc: Document, title: str, subtitle: str):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title); run.bold = True; run.font.size = Pt(20)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(subtitle); run2.font.size = Pt(12)
    doc.add_paragraph("")

def _add_meta_table(doc: Document, meta: dict):
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in meta.items():
        row = t.add_row().cells
        row[0].text = str(k); row[1].text = str(v)

def add_plain_text(doc: Document, txt: str):
    for line in (txt or "").splitlines():
        doc.add_paragraph(line)

def build_word_report(app_name: str, meta: dict, pages: dict[int, str]) -> bytes:
    doc = Document()
    _doc_set_normal_style(doc)
    _add_cover(doc, app_name, "Akademik hisobot (AI tahlil + transliteratsiya + tarjima + izoh)")
    _add_meta_table(doc, meta)
    doc.add_page_break()

    page_keys = sorted(pages.keys())
    for j, idx in enumerate(page_keys):
        doc.add_heading(f"Varaq {idx+1}", level=1)
        add_plain_text(doc, pages[idx] or "")
        if j != len(page_keys) - 1:
            doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==========================================
# 10) SIDEBAR
# ==========================================
with st.sidebar:
    st.markdown("<h2 style='text-align:center;'>📜 MS AI PRO</h2>", unsafe_allow_html=True)

    st.markdown("### ✉️ Email bilan kirish")
    st.caption("Email kiriting — kreditlar va premium funksiyalar ochiladi.")

    email_in = st.text_input("Email", value=(st.session_state.u_email or ""), placeholder="example@mail.com")
    if st.button("KIRISH"):
        email = (email_in or "").strip().lower()
        if not email or "@" not in email:
            st.error("Emailni to‘g‘ri kiriting.")
        else:
            st.session_state.auth = True
            st.session_state.u_email = email
            ensure_profile(email)
            st.rerun()

    if st.session_state.auth:
        st.divider()
        credits = get_credits(st.session_state.u_email)
        st.markdown(f"""
        <div style="
          background: rgba(255,255,255,0.04);
          border: 1px solid rgba(197,160,89,0.35);
          border-radius: 16px;
          padding: 12px 12px;
          box-shadow: 0 14px 30px rgba(0,0,0,0.18);
        ">
          <div style="font-weight:900; color:{C["gold"]}; font-size:14px;">👤 Profil</div>
          <div style="color:{C["text"]}; margin-top:6px; font-weight:900;">{html.escape(st.session_state.u_email)}</div>
          <div style="margin-top:8px; color:{C["muted"]};">Kreditlar: <span style="color:{C["gold"]}; font-weight:900;">{credits}</span> sahifa</div>
          <div style="margin-top:8px; color:{C["muted"]}; font-size:12px;">Premium: Word • Tahrir • Chat • History • Save results</div>
        </div>
        """, unsafe_allow_html=True)

        if st.session_state.warn_db:
            st.warning("DB yoki RPC’da vaqtinchalik muammo bo‘lishi mumkin (fallback ishlayapti).")

        if st.button("🚪 CHIQISH"):
            st.session_state.auth = False
            st.session_state.u_email = ""
            st.rerun()

        with st.expander(f"🧾 History (oxirgi {HISTORY_LIMIT})", expanded=False):
            try:
                r = db.table("usage_logs") \
                    .select("created_at,doc_name,page_index,status") \
                    .eq("email", st.session_state.u_email) \
                    .order("created_at", desc=True) \
                    .limit(HISTORY_LIMIT).execute()
                data = r.data or []
                if not data:
                    st.caption("History hozircha bo‘sh.")
                else:
                    for row in data:
                        ts = (row.get("created_at") or "")[:19].replace("T", " ")
                        st.caption(f"{ts} • {row.get('doc_name','')} • sahifa {int(row.get('page_index',0))+1} • {row.get('status','')}")
            except Exception:
                st.caption("History o‘qilmadi (policy/RLS yoki DB muammo).")

    st.divider()

    st.markdown("### Google bilan kirish")
    st.markdown(f"""
    <div style="
      opacity: 0.55;
      pointer-events: none;
      background: rgba(255,255,255,0.03);
      border: 1px solid rgba(197,160,89,0.30);
      border-radius: 14px;
      padding: 10px 10px;
      color: {C["muted"]};
      font-weight: 900;
      text-align:center;
    ">
      ⏳ Google login vaqtincha texnik ishlar
    </div>
    """, unsafe_allow_html=True)

    st.divider()

    auto_detect = st.checkbox("🧠 Avto aniqlash (tavsiya)", value=True)
    lang = st.selectbox("Taxminiy matn tili (hint):", ["Noma'lum", "Chig'atoy", "Forscha", "Arabcha", "Eski Turkiy"], index=0)
    era = st.selectbox("Taxminiy xat uslubi (hint):", ["Noma'lum", "Nasta'liq", "Suls", "Riq'a", "Kufiy"], index=0)

    st.markdown("### 🧪 Skan sozlamalari")
    rotate = st.select_slider("Aylantirish:", options=[0, 90, 180, 270], value=0)
    brightness = st.slider("Yorqinlik:", 0.5, 2.0, 1.0)
    contrast = st.slider("Kontrast:", 0.5, 3.0, 1.35)
    sharpen = st.slider("Sharpen:", 0.0, 1.5, 0.9, 0.1)

    scale = st.slider("PDF render scale:", 1.5, 3.8, 2.4, 0.1)
    max_pages = st.slider("Preview max sahifa:", 1, 60, 30)

    st.markdown("### 🧭 Ko'rinish")
    view_mode = st.radio("Natija ko'rinishi:", ["Yonma-yon", "Tabs"], index=0, horizontal=True)

    st.divider()
    st.markdown("### ⚙️ Tahlil rejimi")
    analysis_mode = st.radio(
        "Rejim:",
        ["Tez", "To‘liq"],
        index=0,
        horizontal=True,
        key="analysis_mode"
    )
    st.caption("Tez: natija tezroq. To‘liq: ko‘proq chuqurlik (sekinroq).")


# ==========================================
# 11) MAIN
# ==========================================
st.title("📜 Manuscript AI Center")
st.markdown("<p style='text-align:center;'>Qadimiy hujjatlarni yuklang va AI yordamida tahlil qiling.</p>", unsafe_allow_html=True)

uploaded_file = st.file_uploader(
    "Faylni yuklang",
    type=["pdf", "png", "jpg", "jpeg"],
    label_visibility="collapsed"
)

if uploaded_file is None:
    st.markdown(f"""
    <div style="
      background: linear-gradient(180deg, rgba(197,160,89,0.18), rgba(255,255,255,0.04));
      border: 1px solid rgba(197,160,89,0.40);
      border-radius: 18px;
      padding: 18px 18px;
      box-shadow: 0 18px 40px rgba(0,0,0,0.22);
      max-width: 980px;
      margin: 18px auto 0 auto;
    ">
      <div style="display:flex; gap:14px; align-items:center; justify-content:space-between; flex-wrap:wrap;">
        <div>
          <div style="font-size:18px; font-weight:900; color:{C["gold"]};">📜 Manuscript AI — Tez demo</div>
          <div style="color:{C["muted"]}; margin-top:4px;">
            1) PDF/rasm yuklang → 2) Sahifani tanlang → 3) Akademik tahlilni boshlang
          </div>
        </div>
        <div style="
          background: rgba(12,20,33,0.55);
          border: 1px solid rgba(197,160,89,0.35);
          border-radius: 14px;
          padding: 10px 12px;
          color:{C["text"]};
          font-weight:900;
        ">
          ✅ Drag & drop ishlaydi
        </div>
      </div>
      <div style="margin-top:12px; color:{C["muted"]}; font-size:14px; line-height:1.6;">
        * Skan sifatini oshirish uchun sidebar’dagi Yorqinlik/Kontrast/Sharpen’dan foydalaning.
        <br/>* Katta PDF bo‘lsa, “Preview max sahifa”ni 20–30 atrofida qoldiring.
      </div>
    </div>
    """, unsafe_allow_html=True)

if uploaded_file:
    if st.session_state.last_fn != uploaded_file.name:
        with st.spinner("Preparing..."):
            file_bytes = uploaded_file.getvalue()
            if uploaded_file.type == "application/pdf":
                pages = render_pdf_pages_to_bytes(file_bytes, max_pages=max_pages, scale=scale)
            else:
                img = Image.open(io.BytesIO(file_bytes))
                pages = [pil_to_jpeg_bytes(img, quality=92, max_side=3000)]

            st.session_state.page_bytes = pages
            st.session_state.last_fn = uploaded_file.name

            st.session_state.results = {}
            st.session_state.chats = {}
            st.session_state.warn_db = False
            gc.collect()

            if st.session_state.auth and st.session_state.u_email:
                restored = load_reports(st.session_state.u_email, st.session_state.last_fn)
                if restored:
                    st.session_state.results.update(restored)

    processed_pages = [
        preprocess_bytes(b, brightness=brightness, contrast=contrast, rotate=rotate, sharpen=sharpen)
        for b in st.session_state.page_bytes
    ]

    total_pages = len(processed_pages)
    st.caption(f"Yuklandi: **{total_pages}** sahifa (preview limit: {max_pages}).")

    if total_pages <= 30:
        selected_indices = st.multiselect(
            "Sahifalarni tanlang:",
            options=list(range(total_pages)),
            default=[0] if total_pages else [],
            format_func=lambda x: f"{x+1}-sahifa"
        )
    else:
        page_spec = st.text_input("Sahifalar (masalan: 1-5, 9, 12-20):", value="1")
        selected_indices = parse_pages(page_spec, total_pages)

    if not st.session_state.auth:
        if len(selected_indices) > DEMO_LIMIT_PAGES:
            st.warning(f"Demo rejim: maksimal {DEMO_LIMIT_PAGES} sahifa tahlil qilinadi. Premium uchun Email bilan kiring.")
            selected_indices = selected_indices[:DEMO_LIMIT_PAGES]

    if not st.session_state.auth:
        st.markdown(f"""
        <div style="
          background: rgba(255,243,224,1);
          border: 1px solid #ffb74d;
          border-radius: 14px;
          padding: 14px 14px;
          margin: 12px 0 12px 0;
          max-width: 980px;
          margin-left:auto; margin-right:auto;
        ">
          <div style="font-weight:900; color:#e65100; font-size:16px;">
            🔒 Word hisobot • Tahrir • AI Chat • History • Save results — Premium
          </div>
          <div style="margin-top:6px; color:#5a3a00; line-height:1.6;">
            Demo rejimda natijani ko‘rishingiz mumkin (maks {DEMO_LIMIT_PAGES} sahifa). To‘liq funksiyalar uchun Email bilan kiring.
          </div>
        </div>
        """, unsafe_allow_html=True)

    if not st.session_state.results and selected_indices:
        cols = st.columns(min(len(selected_indices), 4))
        for i, idx in enumerate(selected_indices[:16]):
            with cols[i % min(len(cols), 4)]:
                st.image(processed_pages[idx], caption=f"Varaq {idx+1}", use_container_width=True)

    if st.session_state.results:
        q = st.text_input("🔎 Natijalarda qidirish (kalit so‘z):", value="", placeholder="masalan: Muhammad, sana, joy nomi...")
        if q.strip():
            hits = []
            for k, txt in st.session_state.results.items():
                if q.lower() in (txt or "").lower():
                    hits.append(k)
            if hits:
                st.success("Topildi: " + ", ".join([f"{i+1}-sahifa" for i in hits]))
            else:
                st.info("Hozircha topilmadi.")

    # RUN analysis (PATCHED)
    if st.button("✨ AKADEMIK TAHLILNI BOSHLASH"):
        hint_lang = "" if (auto_detect or lang == "Noma'lum") else lang
        hint_era = "" if (auto_detect or era == "Noma'lum") else era
        p_read, p_an = build_prompts(hint_lang, hint_era)

        cfg = get_analysis_cfg(st.session_state.get("analysis_mode", "Tez"))
        min_chars = int(cfg.get("min_chars", 420))

        total = len(selected_indices)
        done = 0
        bar = st.progress(0.0) if total > 0 else None
        ph = st.empty()

        def upd():
            if bar:
                bar.progress(done / max(total, 1))
            ph.markdown(
                f"<div style='text-align:center; color:{C['muted']}; font-weight:900;'>"
                f"📈 Tahlil: <span style='color:{C['gold']};'>{done}/{total}</span>"
                f"</div>",
                unsafe_allow_html=True
            )

        upd()

        for idx in selected_indices:
            time.sleep(random.uniform(*BATCH_DELAY_RANGE))
            reserved = False

            with st.status(f"Sahifa {idx+1}...") as s:
                try:
                    if st.session_state.auth:
                        ok = consume_credit_safe(st.session_state.u_email, 1)
                        if not ok:
                            s.update(label="Kredit yetarli emas", state="error")
                            st.warning("Kredit tugagan. Davom etish uchun kredit qo‘shing.")
                            log_usage(st.session_state.u_email, st.session_state.last_fn, idx, "no_credits")
                            continue
                        reserved = True

                    call_budget = {"n": 0}

                    img_bytes = processed_pages[idx]
                    payload = {"mime_type": "image/jpeg", "data": base64.b64encode(img_bytes).decode("utf-8")}

                    # --- STEP A: READ ---
                    read_text = _call_gemini_guarded(p_read, [payload], call_budget, cfg, tries=3).strip()
                    translit_lines = extract_translit_lines(read_text)
                    diag_block = extract_diag_block(read_text)

                    if (not has_read_sections(read_text)) or (not translit_min_len_ok(translit_lines, min_chars=min_chars)):
                        d2, tl2 = read_transliteration_with_strips(
                            p_read, img_bytes,
                            strips_n=int(cfg.get("strips_default", 6)),
                            min_chars=min_chars,
                            call_budget=call_budget,
                            cfg=cfg
                        )
                        if tl2.strip():
                            diag_block = d2 or diag_block
                            translit_lines = tl2

                    if not translit_min_len_ok(translit_lines, min_chars=min_chars):
                        d3, tl3 = read_transliteration_with_strips(
                            p_read, img_bytes,
                            strips_n=int(cfg.get("strips_retry", 9)),
                            min_chars=min_chars,
                            call_budget=call_budget,
                            cfg=cfg
                        )
                        if tl3.strip():
                            diag_block = d3 or diag_block
                            translit_lines = tl3

                    read_text = build_read_text(diag_block, translit_lines)

                    if not read_text.strip() or not translit_lines.strip():
                        if reserved:
                            refund_credit_safe(st.session_state.u_email, 1)
                        s.update(label="Bo‘sh natija (refund)", state="error")
                        if st.session_state.auth:
                            log_usage(st.session_state.u_email, st.session_state.last_fn, idx, "empty_read")
                        continue

                    # --- STEP B: ANALYZE (2-6) ---
                    try:
                        an_text = build_analyze_text(p_an, translit_lines, call_budget=call_budget, cfg=cfg)
                    except Exception:
                        an_prompt = p_an + "\n\nQuyidagi transliteratsiya bilan ishlang (faqat shunga tayangan holda):\n" + translit_lines
                        an_text = _call_gemini_guarded(an_prompt, [], call_budget, cfg, tries=3).strip()

                    final_text = read_text.strip() + "\n\n" + (an_text.strip() or "")

                    if not has_final_sections(final_text):
                        try:
                            p_an2 = p_an + "\n\nMUHIM: Formatni qat'iy saqla. 2-6 bo'limlarning hammasini to'liq yoz."
                            an_text2 = build_analyze_text(p_an2, translit_lines, call_budget=call_budget, cfg=cfg)
                            if an_text2.strip():
                                final_text = read_text.strip() + "\n\n" + an_text2.strip()
                        except Exception:
                            pass

                    st.session_state.results[idx] = final_text
                    s.update(label="Tayyor!", state="complete")

                    if st.session_state.auth:
                        save_report(st.session_state.u_email, st.session_state.last_fn, idx, final_text)
                        log_usage(st.session_state.u_email, st.session_state.last_fn, idx, "ok")

                except Exception as e:
                    if reserved:
                        refund_credit_safe(st.session_state.u_email, 1)
                    s.update(label="Xato (refund)", state="error")
                    st.error(f"Xato: {e}")
                    if st.session_state.auth:
                        log_usage(st.session_state.u_email, st.session_state.last_fn, idx, "error", note=str(e))

            done += 1
            upd()

        if bar:
            bar.progress(1.0)

        gc.collect()
        st.rerun()

    # RESULTS
    if st.session_state.results:
        st.divider()
        keys = sorted(st.session_state.results.keys())

        jump = st.selectbox("⚡ Tez o‘tish (natija bor sahifalar):", options=keys, format_func=lambda x: f"{x+1}-sahifa")
        keys = [jump] + [k for k in keys if k != jump]

        for idx in keys:
            with st.expander(f"📖 Varaq {idx+1}", expanded=True):
                res = st.session_state.results.get(idx, "") or ""

                img_b64 = base64.b64encode(processed_pages[idx]).decode("utf-8")
                img_html = f"""
                <div class="sticky-preview">
                  <img src="data:image/jpeg;base64,{img_b64}" alt="page {idx+1}" />
                </div>
                """

                copy_js = f"""
                <button id="copybtn" style="
                    width:100%;
                    padding:10px 12px;
                    border-radius:12px;
                    border:1px solid rgba(0,0,0,0.12);
                    font-weight:900;
                    cursor:pointer;
                ">📋 Natijani nusxalash</button>
                <script>
                  const txt = {html.escape(res)!r};
                  document.getElementById("copybtn").onclick = async () => {{
                    try {{
                      await navigator.clipboard.writeText(txt);
                      document.getElementById("copybtn").innerText = "✅ Nusxalandi";
                      setTimeout(()=>document.getElementById("copybtn").innerText="📋 Natijani nusxalash", 1500);
                    }} catch(e) {{
                      document.getElementById("copybtn").innerText = "❌ Clipboard ruxsat yo‘q";
                    }}
                  }}
                </script>
                """

                if view_mode == "Tabs":
                    tabs = st.tabs(["📷 Rasm", "📝 Natija", "✍️ Tahrir", "💬 Chat"])
                    with tabs[0]:
                        st.markdown(img_html, unsafe_allow_html=True)
                    with tabs[1]:
                        components.html(copy_js, height=55)
                        components.html(render_result_card(res, C["gold"]), height=580, scrolling=True)
                    with tabs[2]:
                        if not st.session_state.auth:
                            st.info("🔒 Tahrir premium. Email bilan kiring.")
                        else:
                            st.session_state.results[idx] = st.text_area(
                                f"Tahrir ({idx+1}):",
                                value=st.session_state.results[idx],
                                height=260,
                                key=f"ed_{idx}"
                            )
                            save_report(st.session_state.u_email, st.session_state.last_fn, idx, st.session_state.results[idx])
                    with tabs[3]:
                        if not st.session_state.auth:
                            st.info("🔒 Chat premium. Email bilan kiring.")
                        else:
                            st.session_state.chats.setdefault(idx, [])
                            for ch in st.session_state.chats[idx]:
                                st.markdown(f"<div class='chat-user'><b>S:</b> {html.escape(ch['q'])}</div>", unsafe_allow_html=True)
                                st.markdown(f"<div class='chat-ai'><b>AI:</b> {html.escape(ch['a'])}</div>", unsafe_allow_html=True)
                            user_q = st.text_input("Savol bering:", key=f"q_{idx}")
                            if st.button(f"So'rash ({idx+1})", key=f"btn_{idx}"):
                                if user_q.strip():
                                    with st.spinner("..."):
                                        chat_prompt = f"Matn: {st.session_state.results[idx]}\nSavol: {user_q}\nJavobni o‘zbekcha, aniq va qisqa yoz."
                                        chat_resp = model.generate_content([chat_prompt], generation_config={"max_output_tokens": 1200, "temperature": 0.2})
                                        st.session_state.chats[idx].append({"q": user_q, "a": getattr(chat_resp, "text", "") or ""})
                                        st.rerun()
                else:
                    c1, c2 = st.columns([1, 1.35], gap="large")
                    with c1:
                        st.markdown(img_html, unsafe_allow_html=True)
                    with c2:
                        components.html(copy_js, height=55)
                        components.html(render_result_card(res, C["gold"]), height=520, scrolling=True)

                        if not st.session_state.auth:
                            st.info("🔒 Word/Tahrir/Chat premium. Email bilan kiring.")
                        else:
                            st.session_state.results[idx] = st.text_area(
                                f"Tahrir ({idx+1}):",
                                value=st.session_state.results[idx],
                                height=220,
                                key=f"ed_{idx}"
                            )
                            save_report(st.session_state.u_email, st.session_state.last_fn, idx, st.session_state.results[idx])

                            with st.expander("💬 AI Chat (shu varaq bo‘yicha)", expanded=True):
                                st.session_state.chats.setdefault(idx, [])
                                for ch in st.session_state.chats[idx]:
                                    st.markdown(f"<div class='chat-user'><b>S:</b> {html.escape(ch['q'])}</div>", unsafe_allow_html=True)
                                    st.markdown(f"<div class='chat-ai'><b>AI:</b> {html.escape(ch['a'])}</div>", unsafe_allow_html=True)

                                user_q = st.text_input("Savol bering:", key=f"q_side_{idx}")
                                if st.button(f"So'rash (Varaq {idx+1})", key=f"btn_side_{idx}"):
                                    if user_q.strip():
                                        with st.spinner("..."):
                                            chat_prompt = f"Matn: {st.session_state.results[idx]}\nSavol: {user_q}\nJavobni o‘zbekcha, aniq va qisqa yoz."
                                            chat_resp = model.generate_content([chat_prompt], generation_config={"max_output_tokens": 1200, "temperature": 0.2})
                                            st.session_state.chats[idx].append({"q": user_q, "a": getattr(chat_resp, "text", "") or ""})
                                            st.rerun()

        if st.session_state.auth and st.session_state.results:
            detected = aggregate_detected_meta(st.session_state.results)
            meta = {
                "Hujjat nomi": st.session_state.last_fn,
                "Til (aniqlangan)": detected["til"] or "Noma'lum",
                "Xat uslubi (aniqlangan)": detected["xat"] or "Noma'lum",
                "Avto aniqlash": "Ha" if auto_detect else "Yo‘q",
                "Til (hint)": lang,
                "Xat uslubi (hint)": era,
                "Eksport qilingan sahifalar": ", ".join(str(i+1) for i in sorted(st.session_state.results.keys())),
                "Yaratilgan vaqt": datetime.now().strftime("%Y-%m-%d %H:%M"),
            }
            report_bytes = build_word_report("Manuscript AI", meta, st.session_state.results)
            st.download_button(
                "📥 WORD HISOBOTNI YUKLAB OLISH (.docx)",
                report_bytes,
                file_name="Manuscript_AI_Report.docx"
            )

gc.collect()
